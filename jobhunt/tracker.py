"""Stage 3: persistent Drive job tracker (Part B).

Runs as `python -m jobhunt.tracker`. Builds one "Houndstooth Job Tracker.xlsx" in the
Drive delivery folder, accumulating every job ever digested (per digest_log), refreshed
daily update-in-place, with user-owned Status/Notes preserved across refreshes.

Reuses `_drive_service()` from jobhunt.deliver -- never edits that module.

CLI:
    python -m jobhunt.tracker [--local-only] [--out PATH] [--db PATH]

--local-only skips the Drive round-trip (no Status/Notes merge, no upload) and writes
only to a local xlsx path. --out defaults to a tempfile path (never reports/, which CI
commits). --db opens an arbitrary sqlite file instead of jobhunt.db.connect() (used for
scratch-copy verification).
"""
import argparse
import csv
import io
import json
import os
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import db
from .deliver import _drive_service

ROOT = Path(__file__).resolve().parent.parent
CONFIG_PATH = ROOT / "config.yaml"

STATUS_CHOICES = ["Interested", "Applied", "Interviewing", "Offer", "Rejected", "Dismissed"]

JOBS_HEADER = ["Score", "Title", "Company", "Location", "First Seen", "Last Seen",
               "Times Seen", "In Latest Digest", "Resume", "Cover Letter", "Job Link",
               "Status", "Notes"]
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
TRACKER_NAME_CONTAINS = "Houndstooth Job Tracker"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GSHEET_MIME = "application/vnd.google-apps.spreadsheet"


# ---------------- dedup key (contract 3) ----------------

def dedup_key(url: str, title: str, company: str) -> str:
    url = (url or "").strip()
    if url:
        return url
    return f"NOLINK::{title}::{company}"


# ---------------- Build rows ----------------

def _fetch_job_digest_rows(con):
    """One row per job_id that has ever appeared in digest_log, joined with its jobs data
    and per-job digest aggregates (first/last/times seen)."""
    return con.execute(
        """
        SELECT j.id, j.score, j.title, j.company, j.location, j.url, j.drive_links,
               j.user_status, j.user_notes,
               d.first_seen, d.last_seen, d.times_seen
        FROM jobs j
        JOIN (
            SELECT job_id, MIN(digest_date) AS first_seen, MAX(digest_date) AS last_seen,
                   COUNT(*) AS times_seen
            FROM digest_log
            GROUP BY job_id
        ) d ON d.job_id = j.id
        """
    ).fetchall()


def build_rows(con):
    """Return (rows, stats). rows sorted score desc, company, title.

    Per dedup-key group (contract 3): core fields (score/title/company/location/url) from
    the constituent job_id row with the latest last_seen; resume/cover from the
    constituent row with non-null drive_links and the latest last_seen among those;
    Status/Notes similarly from the latest non-empty among the group. first_seen/last_seen
    aggregate across the group (min/max); times_seen sums across the group (the
    url-collapse quirk: rare cross-listed same-url rows, documented on the About sheet).
    in_latest = last_seen == MAX(digest_date) over the whole table.
    """
    job_rows = _fetch_job_digest_rows(con)
    overall_max = con.execute("SELECT MAX(digest_date) AS d FROM digest_log").fetchone()["d"]

    groups: dict[str, list] = {}
    for r in job_rows:
        key = dedup_key(r["url"], r["title"], r["company"])
        groups.setdefault(key, []).append(r)

    rows = []
    for key, members in groups.items():
        core = max(members, key=lambda r: r["last_seen"])
        first_seen = min(m["first_seen"] for m in members)
        last_seen = max(m["last_seen"] for m in members)
        times_seen = sum(m["times_seen"] for m in members)

        drive_candidates = [m for m in members if m["drive_links"]]
        drive_src = max(drive_candidates, key=lambda r: r["last_seen"]) if drive_candidates else None

        status_candidates = [m for m in members if m["user_status"]]
        status_src = max(status_candidates, key=lambda r: r["last_seen"]) if status_candidates else None
        notes_candidates = [m for m in members if m["user_notes"]]
        notes_src = max(notes_candidates, key=lambda r: r["last_seen"]) if notes_candidates else None

        resume_url = cover_url = None
        if drive_src and drive_src["drive_links"]:
            links = json.loads(drive_src["drive_links"])
            resume_url = links.get("resume")
            cover_url = links.get("cover_letter")

        rows.append({
            "dedup_key": key,
            "job_ids": [m["id"] for m in members],
            "score": core["score"],
            "title": core["title"],
            "company": core["company"],
            "location": core["location"],
            "url": core["url"],
            "first_seen": first_seen,
            "last_seen": last_seen,
            "times_seen": times_seen,
            "in_latest": last_seen == overall_max,
            "resume_url": resume_url,
            "cover_url": cover_url,
            "status": status_src["user_status"] if status_src else "",
            "notes": notes_src["user_notes"] if notes_src else "",
        })

    rows.sort(key=lambda r: (-(r["score"] or -1), r["company"] or "", r["title"] or ""))

    stats = {
        "total_unique": len(rows),
        "in_latest": sum(1 for r in rows if r["in_latest"]),
    }
    return rows, stats


# ---------------- Status/Notes round-trip (skipped under --local-only) ----------------

def find_existing_tracker(service, folder_id: str):
    """Most-recently-modified non-trashed file in folder_id whose name contains the
    tracker name. Returns the files() dict, or None."""
    q = (f"name contains '{TRACKER_NAME_CONTAINS}' and '{folder_id}' in parents "
         f"and trashed=false")
    res = service.files().list(q=q, fields="files(id,name,mimeType,modifiedTime)",
                                orderBy="modifiedTime desc", pageSize=10).execute()
    files = res.get("files", [])
    return files[0] if files else None


def _parse_status_notes_rows(rows_iter):
    """rows_iter yields lists aligned to JOBS_HEADER (or a CSV export of it, same order).
    Returns {dedup_key: (status, notes)} keyed by Job Link, falling back to the NOLINK key
    built from Title+Company when Job Link is blank."""
    header = None
    idx = {}
    out = {}
    for row in rows_iter:
        if header is None:
            header = [str(c or "").strip() for c in row]
            for name in ("Title", "Company", "Job Link", "Status", "Notes"):
                idx[name] = header.index(name)
            continue
        if not row or all(c in (None, "") for c in row):
            continue
        title = row[idx["Title"]] if idx["Title"] < len(row) else ""
        company = row[idx["Company"]] if idx["Company"] < len(row) else ""
        job_link = row[idx["Job Link"]] if idx["Job Link"] < len(row) else ""
        status = row[idx["Status"]] if idx["Status"] < len(row) else ""
        notes = row[idx["Notes"]] if idx["Notes"] < len(row) else ""
        key = dedup_key(job_link, title, company)
        out[key] = (status or "", notes or "")
    return out


def read_status_notes(service, file_meta) -> dict:
    """Read {dedup_key: (status, notes)} from the existing tracker file: native Google
    Sheet -> CSV export; real xlsx -> get_media bytes + openpyxl."""
    file_id = file_meta["id"]
    if file_meta.get("mimeType") == GSHEET_MIME:
        data = service.files().export(fileId=file_id, mimeType="text/csv").execute()
        text = data.decode("utf-8") if isinstance(data, bytes) else data
        reader = csv.reader(io.StringIO(text))
        return _parse_status_notes_rows(reader)
    else:
        data = service.files().get_media(fileId=file_id).execute()
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb["Jobs"] if "Jobs" in wb.sheetnames else wb.active
        return _parse_status_notes_rows(ws.iter_rows(values_only=True))


def safe_read_status_notes(service, file_meta) -> dict:
    """read_status_notes, degraded: a corrupt/unreadable existing tracker returns {} with
    a warning instead of killing the refresh -- the subsequent update-in-place overwrites
    the bad file with a freshly built workbook. (Observed live 2026-08-07: the prior
    cowork run left a truncated xlsx whose zip central directory pointed past EOF.)"""
    try:
        return read_status_notes(service, file_meta)
    except Exception as e:
        print(f"[tracker] WARNING: could not read existing tracker "
              f"'{file_meta.get('name', file_meta.get('id'))}': {e}; "
              "proceeding without Status/Notes merge")
        return {}


def merge_into_db(con, rows, status_notes_map: dict):
    """Merge round-tripped Status/Notes into jobs.user_status/user_notes (contract 5):
    non-empty incoming wins; empty incoming never clears a non-empty stored value.
    Applies to every constituent job_id in a dedup group so a future rebuild sees it
    regardless of which job_id becomes the "latest" row next time."""
    for r in rows:
        incoming = status_notes_map.get(r["dedup_key"])
        if not incoming:
            continue
        inc_status, inc_notes = incoming
        for job_id in r["job_ids"]:
            if inc_status:
                con.execute("UPDATE jobs SET user_status=? WHERE id=?", (inc_status, job_id))
            if inc_notes:
                con.execute("UPDATE jobs SET user_notes=? WHERE id=?", (inc_notes, job_id))
    con.commit()


# ---------------- Write xlsx ----------------

def write_xlsx(rows, stats, out_path, summary_min_score: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(JOBS_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for r in rows:
        row_values = [
            r["score"], r["title"], r["company"], r["location"],
            r["first_seen"], r["last_seen"], r["times_seen"],
            "TRUE" if r["in_latest"] else "FALSE",
            r["resume_url"] or "—", r["cover_url"] or "—",
            r["url"] or "", r["status"] or "", r["notes"] or "",
        ]
        ws.append(row_values)
        excel_row = ws.max_row

        title_cell = ws.cell(row=excel_row, column=2)
        if r["url"]:
            title_cell.hyperlink = r["url"]
            title_cell.style = "Hyperlink"

        resume_cell = ws.cell(row=excel_row, column=9)
        if r["resume_url"]:
            resume_cell.hyperlink = r["resume_url"]
            resume_cell.style = "Hyperlink"

        cover_cell = ws.cell(row=excel_row, column=10)
        if r["cover_url"]:
            cover_cell.hyperlink = r["cover_url"]
            cover_cell.style = "Hyperlink"

        if (r["score"] or 0) >= summary_min_score:
            for col in range(1, len(JOBS_HEADER) + 1):
                ws.cell(row=excel_row, column=col).fill = GREEN_FILL

    last_row = ws.max_row
    last_col_letter = get_column_letter(len(JOBS_HEADER))
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_CHOICES) + '"',
                         allow_blank=True)
    ws.add_data_validation(dv)
    if last_row >= 2:
        dv.add(f"L2:L{last_row}")

    ws.column_dimensions["B"].width = 45  # Title
    ws.column_dimensions["K"].width = 60  # Job Link
    ws.column_dimensions["C"].width = 28  # Company
    ws.column_dimensions["M"].width = 30  # Notes

    about = wb.create_sheet("About")
    stamp = datetime.now(ZoneInfo("America/Los_Angeles")).strftime("%Y-%m-%d %H:%M:%S %Z")
    about_lines = [
        ("Generated", stamp),
        ("Total unique jobs", stats["total_unique"]),
        ("In latest digest", stats["in_latest"]),
        ("", ""),
        ("Note", "Status and Notes columns are user-owned. Refreshes preserve your edits: "
                 "an empty Status/Notes cell here never overwrites a value you've already set."),
        ("Note", "Jobs cross-listed under the same URL collapse into a single row here "
                 "(the dedup key is the job URL). This affects roughly 5 of ~2,744 rows."),
    ]
    for row in about_lines:
        about.append(row)
    for cell in about["A"]:
        cell.font = Font(bold=True)

    wb.save(out_path)
    return out_path


def sanity_check(out_path, rows):
    """Read-back sanity check before any upload. Raises on mismatch."""
    wb = load_workbook(out_path)
    if "Jobs" not in wb.sheetnames:
        raise RuntimeError(f"sanity check failed: no 'Jobs' sheet in {out_path}")
    ws = wb["Jobs"]
    header = [c.value for c in ws[1]]
    if header != JOBS_HEADER:
        raise RuntimeError(f"sanity check failed: header mismatch {header} != {JOBS_HEADER}")
    data_row_count = ws.max_row - 1
    if data_row_count != len(rows):
        raise RuntimeError(f"sanity check failed: row count {data_row_count} != expected {len(rows)}")
    return data_row_count


# ---------------- Upload (contract 6) ----------------

def upload_tracker(service, folder_id: str, out_path, tracker_filename: str, existing):
    """existing: files().list dict or None (from find_existing_tracker).
    Real xlsx -> files().update on the same fileId. No file, or existing is a native
    Google Sheet -> files().create + warning naming the stale Sheet."""
    from googleapiclient.http import MediaFileUpload

    media = MediaFileUpload(str(out_path), mimetype=XLSX_MIME)
    warning = None

    if existing and existing.get("mimeType") != GSHEET_MIME:
        updated = service.files().update(fileId=existing["id"], media_body=media,
                                          fields="id, webViewLink").execute()
        return updated["webViewLink"], warning

    if existing and existing.get("mimeType") == GSHEET_MIME:
        warning = (f"Existing '{existing['name']}' (id={existing['id']}) is a native Google "
                   f"Sheet and can't be media-updated to xlsx; created a new xlsx tracker "
                   f"instead. Remove the stale Sheet manually.")

    meta = {"name": tracker_filename, "parents": [folder_id]}
    created = service.files().create(body=meta, media_body=media,
                                      fields="id, webViewLink").execute()
    return created["webViewLink"], warning


# ---------------- Orchestration ----------------

def _load_config(config_path=CONFIG_PATH):
    return yaml.safe_load(open(config_path)).get("delivery", {})


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--local-only", action="store_true",
                         help="skip the Drive Status/Notes round-trip and upload")
    parser.add_argument("--out", default=None, help="xlsx output path (default: tempfile)")
    parser.add_argument("--db", default=None, help="sqlite db path (default: jobhunt.db.connect())")
    return parser.parse_args(argv)


def run(argv=None):
    args = parse_args(argv)
    cfg = _load_config()
    summary_min_score = cfg.get("summary_min_score", 78)
    tracker_filename = cfg.get("tracker_file", "Houndstooth Job Tracker.xlsx")

    if args.db:
        con = sqlite3.connect(args.db)
        con.row_factory = sqlite3.Row
    else:
        con = db.connect()

    rows, stats = build_rows(con)

    service = None
    existing = None
    if not args.local_only:
        folder_id = cfg.get("gdrive_folder_id") or os.environ.get("GDRIVE_FOLDER_ID", "")
        if not folder_id:
            raise RuntimeError("set delivery.gdrive_folder_id in config.yaml or GDRIVE_FOLDER_ID env var")
        service = _drive_service()
        existing = find_existing_tracker(service, folder_id)
        if existing:
            status_notes_map = safe_read_status_notes(service, existing)
            if status_notes_map:
                merge_into_db(con, rows, status_notes_map)
                rows, stats = build_rows(con)  # rebuild so the sheet reflects the merge

    out_path = args.out or str(Path(tempfile.gettempdir()) / f"houndstooth_tracker_{os.getpid()}.xlsx")
    write_xlsx(rows, stats, out_path, summary_min_score)
    row_count = sanity_check(out_path, rows)

    if args.local_only:
        print(f"[tracker] local-only: wrote {out_path}")
    else:
        folder_id = cfg.get("gdrive_folder_id") or os.environ.get("GDRIVE_FOLDER_ID", "")
        view_url, upload_warning = upload_tracker(service, folder_id, out_path, tracker_filename, existing)
        if upload_warning:
            print(f"[tracker] WARNING: {upload_warning}")
        print(f"[tracker] uploaded: {view_url}")

    print(f"[tracker] rows: {row_count} total unique jobs, {stats['in_latest']} in latest digest")
    return {"out_path": out_path, "row_count": row_count, "stats": stats}


if __name__ == "__main__":
    run()
