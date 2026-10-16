"""Stage C: verify a rendered document is real, readable, ATS-parseable text, and does not
read as machine output (the slop detector from docs/plan-2026-10/09-design-principle-human-first.md).

PDF path uses pypdf; DOCX path reads paragraphs via python-docx (table cells included, since
the header is a borderless table).
"""
import json
import re
from pathlib import Path

MIN_CHARS = 200
_MOJIBAKE_RE = re.compile(r"[�\x00-\x08\x0b\x0c\x0e-\x1f]")

BANNED_PHRASES = (
    "results-driven", "passionate", "leverage", "leveraging", "cutting-edge", "synergy",
    "proven track record", "forward-deployed engineer with", "seamless", "robust solutions",
    "fast-paced", "i am excited to", "i am confident that", "game-changer", "best-in-class",
    "world-class", "spearheaded",
)
MAX_EM_DASHES = 2
UNIFORM_TOLERANCE = 0.15   # bullet lengths within this fraction of the mean, 4+ bullets -> uniform
UNIFORM_MIN_BULLETS = 4


def extract_text(doc_path: Path) -> str:
    if doc_path.suffix == ".docx":
        from docx import Document
        doc = Document(str(doc_path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    from pypdf import PdfReader
    reader = PdfReader(str(doc_path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def verify_document(doc_path: Path, required_keywords: list[str]) -> dict:
    text = extract_text(doc_path)
    reasons = []
    if len(text.strip()) < MIN_CHARS:
        reasons.append(f"extracted text too short ({len(text.strip())} chars < {MIN_CHARS})")
    hits = _MOJIBAKE_RE.findall(text)
    if hits:
        reasons.append(f"found {len(hits)} mojibake/control characters in extracted text")
    low = text.lower()
    missing = [kw for kw in required_keywords if kw.lower() not in low]
    if missing:
        reasons.append(f"missing mirrored keywords: {', '.join(missing)}")
    return {"file": str(doc_path), "pass": not reasons, "reasons": reasons,
            "char_count": len(text.strip()), "missing_keywords": missing}


# ---------------------------------------------------------------- slop detector

def _ngrams(text: str, n: int = 4) -> set:
    words = re.findall(r"[a-z0-9']+", text.lower())
    return {" ".join(words[i:i + n]) for i in range(len(words) - n + 1)}


def _uniform(lengths: list[int], minimum: int = UNIFORM_MIN_BULLETS) -> bool:
    if len(lengths) < minimum:
        return False
    mean = sum(lengths) / len(lengths)
    return all(abs(x - mean) <= UNIFORM_TOLERANCE * mean for x in lengths)


def slop_reasons(resume_text: str, cover_text: str, draft: dict | None = None,
                 facts: dict | None = None, cover_paragraphs: list[str] | None = None) -> list[str]:
    reasons = []
    whole = (resume_text + "\n" + cover_text).lower()
    banned = sorted({p for p in BANNED_PHRASES if p in whole})
    if banned:
        reasons.append(f"banned phrases present: {', '.join(banned)}")
    dashes = whole.count("—")
    if dashes > MAX_EM_DASHES:
        reasons.append(f"{dashes} em dashes (max {MAX_EM_DASHES})")

    # Per-section bullet uniformity, from the draft so we know section membership.
    if draft and facts:
        selected = draft.get("resume", {}).get("selected_bullets", [])
        texts = {b["id"]: b["text"] for e in facts.get("experience", []) for b in e.get("bullets", [])}
        texts.update({b["id"]: b["text"] for p in facts.get("projects", []) for b in p.get("bullets", [])})
        for exp in facts.get("experience", []):
            ids = [b["id"] for b in exp.get("bullets", []) if b["id"] in selected]
            if _uniform([len(texts[i]) for i in ids]):
                reasons.append(f"uniform bullet lengths under {exp.get('title', '')}")
        # 4-gram overlap between summary and cover letter
        summary = ""
        if "summary_id" in draft.get("resume", {}):
            for s in facts.get("voice_bank", {}).get("summaries", []):
                if s["id"] == draft["resume"]["summary_id"]:
                    summary = s["text"]
        else:
            summary = draft.get("resume", {}).get("summary", "")
        shared = _ngrams(summary) & _ngrams(cover_text)
        if len(shared) >= 3:
            reasons.append(f"summary and cover letter share {len(shared)} four-word phrases")

    if cover_paragraphs:
        body = [p for p in cover_paragraphs if len(p.split()) > 12]
        lengths = [len(p.split()) for p in body]
        if _uniform(lengths, minimum=3):
            reasons.append("cover letter paragraphs are near-equal length")
        words = sum(lengths)
        if words and not 180 <= words <= 420:
            reasons.append(f"cover letter body is {words} words (want 180-420)")
    return reasons


def verify_job(resume_doc: Path, cover_doc: Path, keywords: list[str],
               draft: dict | None = None, facts: dict | None = None,
               cover_paragraphs: list[str] | None = None) -> dict:
    resume_result = verify_document(resume_doc, keywords)
    cover_result = verify_document(cover_doc, [])
    resume_text = extract_text(resume_doc)
    cover_text = extract_text(cover_doc)
    slop = slop_reasons(resume_text, cover_text, draft, facts, cover_paragraphs)
    return {
        "pass": resume_result["pass"] and cover_result["pass"] and not slop,
        "resume": resume_result,
        "cover_letter": cover_result,
        "slop": slop,
    }


def write_verification(result: dict, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, indent=2))
    return out_path
