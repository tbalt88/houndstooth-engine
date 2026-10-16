"""Offline tests for jobhunt/tracker.py: row build, dedup, round-trip merge semantics,
xlsx writing, and the mocked Drive upload branch logic. No network, no real jobs.db.

Run: .venv\\Scripts\\python.exe -m pytest tests/test_tracker.py -v
"""
import io
import sys
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from jobhunt import db, tracker  # noqa: E402

JOB_COLS = ("id", "source", "source_rank", "title", "company", "location", "url",
            "score", "drive_links", "user_status", "user_notes")


def _make_db(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    return db.connect()


def _insert_job(con, **kwargs):
    row = {c: None for c in JOB_COLS}
    row.update(kwargs)
    con.execute(
        f"INSERT INTO jobs ({', '.join(JOB_COLS)}) VALUES ({', '.join('?' * len(JOB_COLS))})",
        tuple(row[c] for c in JOB_COLS),
    )
    con.commit()


def _log(con, job_id, *dates):
    for d in dates:
        con.execute("INSERT OR IGNORE INTO digest_log (job_id, digest_date) VALUES (?, ?)", (job_id, d))
    con.commit()


# ---------------- Row build: aggregates, dedup, in_latest ----------------

def test_build_rows_aggregates_first_last_times_seen(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng", company="Acme", location="Remote",
                url="https://x/1", score=80)
    _log(con, "j1", "2026-08-01", "2026-08-03", "2026-08-05")

    rows, stats = tracker.build_rows(con)
    assert len(rows) == 1
    r = rows[0]
    assert r["first_seen"] == "2026-08-01"
    assert r["last_seen"] == "2026-08-05"
    assert r["times_seen"] == 3
    assert stats["total_unique"] == 1


def test_build_rows_dedup_collapses_same_url_rows_and_sums(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    # two different job_ids (cross-listed), same url
    _insert_job(con, id="j1", title="Eng A", company="Acme A", location="SD",
                url="https://x/same", score=70)
    _insert_job(con, id="j2", title="Eng B", company="Acme B", location="SF",
                url="https://x/same", score=90)
    _log(con, "j1", "2026-08-01")
    _log(con, "j2", "2026-08-02")

    rows, stats = tracker.build_rows(con)
    assert len(rows) == 1
    r = rows[0]
    assert stats["total_unique"] == 1
    assert r["first_seen"] == "2026-08-01"
    assert r["last_seen"] == "2026-08-02"
    assert r["times_seen"] == 2  # summed across the collapsed group
    assert set(r["job_ids"]) == {"j1", "j2"}


def test_build_rows_core_fields_from_latest_last_seen_row(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Old Title", company="Old Co", location="SD",
                url="https://x/same", score=50)
    _insert_job(con, id="j2", title="New Title", company="New Co", location="SF",
                url="https://x/same", score=95)
    _log(con, "j1", "2026-08-01")
    _log(con, "j2", "2026-08-04")  # later -> wins for core fields

    rows, _ = tracker.build_rows(con)
    r = rows[0]
    assert r["title"] == "New Title"
    assert r["company"] == "New Co"
    assert r["score"] == 95


def test_build_rows_latest_non_null_drive_links_wins(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng", company="Acme", location="SD",
                url="https://x/same", score=70,
                drive_links='{"resume": "https://drive/r1", "cover_letter": "https://drive/c1"}')
    _insert_job(con, id="j2", title="Eng", company="Acme", location="SD",
                url="https://x/same", score=70, drive_links=None)
    _log(con, "j1", "2026-08-01")
    _log(con, "j2", "2026-08-03")  # later appearance, but no drive_links -> j1's link should still surface

    rows, _ = tracker.build_rows(con)
    r = rows[0]
    assert r["resume_url"] == "https://drive/r1"
    assert r["cover_url"] == "https://drive/c1"


def test_build_rows_nolink_key_does_not_collapse_distinct_empty_url_jobs(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng One", company="Acme", location="SD", url="", score=70)
    _insert_job(con, id="j2", title="Eng Two", company="Beta", location="SF", url=None, score=60)
    _log(con, "j1", "2026-08-01")
    _log(con, "j2", "2026-08-01")

    rows, _ = tracker.build_rows(con)
    assert len(rows) == 2
    keys = {r["dedup_key"] for r in rows}
    assert keys == {"NOLINK::Eng One::Acme", "NOLINK::Eng Two::Beta"}


def test_build_rows_in_latest_flag(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Fresh", company="Acme", location="SD", url="https://x/1", score=80)
    _insert_job(con, id="j2", title="Stale", company="Beta", location="SF", url="https://x/2", score=80)
    _log(con, "j1", "2026-08-05")
    _log(con, "j2", "2026-08-01")

    rows, stats = tracker.build_rows(con)
    by_title = {r["title"]: r for r in rows}
    assert by_title["Fresh"]["in_latest"] is True
    assert by_title["Stale"]["in_latest"] is False
    assert stats["in_latest"] == 1


# ---------------- Merge semantics (contract 5) ----------------

def test_merge_non_empty_incoming_wins(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng", company="Acme", location="SD", url="https://x/1", score=70)
    rows = [{"dedup_key": "https://x/1", "job_ids": ["j1"]}]
    tracker.merge_into_db(con, rows, {"https://x/1": ("Applied", "great fit")})

    row = con.execute("SELECT user_status, user_notes FROM jobs WHERE id='j1'").fetchone()
    assert row["user_status"] == "Applied"
    assert row["user_notes"] == "great fit"


def test_merge_empty_incoming_never_clears_stored_value(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng", company="Acme", location="SD", url="https://x/1",
                score=70, user_status="Interviewing", user_notes="phone screen done")
    rows = [{"dedup_key": "https://x/1", "job_ids": ["j1"]}]
    tracker.merge_into_db(con, rows, {"https://x/1": ("", "")})

    row = con.execute("SELECT user_status, user_notes FROM jobs WHERE id='j1'").fetchone()
    assert row["user_status"] == "Interviewing"
    assert row["user_notes"] == "phone screen done"


def test_merge_no_incoming_entry_leaves_row_untouched(tmp_path, monkeypatch):
    con = _make_db(tmp_path, monkeypatch)
    _insert_job(con, id="j1", title="Eng", company="Acme", location="SD", url="https://x/1",
                score=70, user_status="Offer", user_notes="signed")
    rows = [{"dedup_key": "https://x/1", "job_ids": ["j1"]}]
    tracker.merge_into_db(con, rows, {})  # no entry for this key at all

    row = con.execute("SELECT user_status, user_notes FROM jobs WHERE id='j1'").fetchone()
    assert row["user_status"] == "Offer"
    assert row["user_notes"] == "signed"


# ---------------- Round-trip parser: CSV export shape + xlsx-bytes shape ----------------

def test_parse_status_notes_rows_csv_shape():
    rows_iter = [
        tracker.JOBS_HEADER,
        [80, "Eng", "Acme", "SD", "2026-08-01", "2026-08-01", 1, "TRUE", "—", "—",
         "https://x/1", "Applied", "good fit"],
        [70, "Eng2", "Beta", "SF", "2026-08-01", "2026-08-01", 1, "FALSE", "—", "—",
         "", "Rejected", "no reply"],
    ]
    parsed = tracker._parse_status_notes_rows(rows_iter)
    assert parsed["https://x/1"] == ("Applied", "good fit")
    assert parsed["NOLINK::Eng2::Beta"] == ("Rejected", "no reply")


def test_parse_status_notes_rows_xlsx_bytes_shape():
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(tracker.JOBS_HEADER)
    ws.append([90, "Eng3", "Gamma", "NY", "2026-08-02", "2026-08-02", 1, "TRUE", "—", "—",
               "https://x/3", "Interviewing", "onsite scheduled"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    loaded = load_workbook(buf)
    parsed = tracker._parse_status_notes_rows(loaded["Jobs"].iter_rows(values_only=True))
    assert parsed["https://x/3"] == ("Interviewing", "onsite scheduled")


def test_safe_read_status_notes_corrupt_xlsx_returns_empty(capsys):
    # Truncated zip: PK magic but a central directory pointing past EOF (the shape the
    # failed cowork upload left on Drive, 2026-08-07) -- must degrade, not raise.
    service = MagicMock()
    service.files().get_media().execute.return_value = b"PK\x03\x04truncated-not-a-zip"
    out = tracker.safe_read_status_notes(
        service, {"id": "f1", "name": "Houndstooth Job Tracker.xlsx", "mimeType": tracker.XLSX_MIME})
    assert out == {}
    assert "WARNING" in capsys.readouterr().out


# ---------------- Xlsx write + read-back sanity ----------------

def _sample_rows():
    return [
        {"dedup_key": "https://x/1", "job_ids": ["j1"], "score": 95, "title": "Staff Eng",
         "company": "Acme", "location": "SD", "url": "https://x/1", "first_seen": "2026-08-01",
         "last_seen": "2026-08-05", "times_seen": 3, "in_latest": True,
         "resume_url": "https://drive/r1", "cover_url": "https://drive/c1",
         "status": "Applied", "notes": "great"},
        {"dedup_key": "NOLINK::Eng::Beta", "job_ids": ["j2"], "score": 60, "title": "Eng",
         "company": "Beta", "location": "SF", "url": "", "first_seen": "2026-08-02",
         "last_seen": "2026-08-02", "times_seen": 1, "in_latest": False,
         "resume_url": None, "cover_url": None, "status": "", "notes": ""},
    ]


def test_write_xlsx_and_sanity_check_roundtrip(tmp_path):
    rows = _sample_rows()
    stats = {"total_unique": len(rows), "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    row_count = tracker.sanity_check(out_path, rows)
    assert row_count == 2

    wb = load_workbook(out_path)
    ws = wb["Jobs"]
    header = [c.value for c in ws[1]]
    assert header == tracker.JOBS_HEADER


def test_write_xlsx_dropdown_validation_present(tmp_path):
    rows = _sample_rows()
    stats = {"total_unique": len(rows), "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    wb = load_workbook(out_path)
    ws = wb["Jobs"]
    dvs = list(ws.data_validations.dataValidation)
    assert len(dvs) == 1
    dv = dvs[0]
    for choice in tracker.STATUS_CHOICES:
        assert choice in dv.formula1
    assert dv.allow_blank is True
    assert "L2:L3" in str(dv.sqref)


def test_write_xlsx_green_fill_on_rows_above_threshold(tmp_path):
    rows = _sample_rows()  # row1 score=95 (>=78), row2 score=60 (<78)
    stats = {"total_unique": len(rows), "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    wb = load_workbook(out_path)
    ws = wb["Jobs"]
    assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith("C6EFCE")
    row3_fill = ws.cell(row=3, column=1).fill.fgColor.rgb
    assert row3_fill in (None, "00000000")


def test_write_xlsx_hyperlinks_set(tmp_path):
    rows = _sample_rows()
    stats = {"total_unique": len(rows), "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    wb = load_workbook(out_path)
    ws = wb["Jobs"]
    assert ws.cell(row=2, column=2).hyperlink.target == "https://x/1"
    assert ws.cell(row=2, column=9).hyperlink.target == "https://drive/r1"
    assert ws.cell(row=2, column=10).hyperlink.target == "https://drive/c1"
    assert ws.cell(row=3, column=2).hyperlink is None  # no url for row 2
    assert ws.cell(row=3, column=9).value == "—"


def test_write_xlsx_about_sheet_numbers_match_stats(tmp_path):
    rows = _sample_rows()
    stats = {"total_unique": 2, "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    wb = load_workbook(out_path)
    about = wb["About"]
    values = {row[0]: row[1] for row in about.iter_rows(values_only=True) if row[0]}
    assert values["Total unique jobs"] == 2
    assert values["In latest digest"] == 1


def test_sanity_check_raises_on_row_count_mismatch(tmp_path):
    rows = _sample_rows()
    stats = {"total_unique": len(rows), "in_latest": 1}
    out_path = tmp_path / "tracker.xlsx"
    tracker.write_xlsx(rows, stats, out_path, summary_min_score=78)

    with pytest.raises(RuntimeError):
        tracker.sanity_check(out_path, rows[:1])  # expect 1, actual file has 2


# ---------------- Upload branch logic (mocked Drive service) ----------------

def test_upload_existing_real_xlsx_updates_same_file_id(tmp_path):
    out_path = tmp_path / "t.xlsx"
    Workbook().save(out_path)
    service = MagicMock()
    service.files.return_value.update.return_value.execute.return_value = {
        "id": "existing-id", "webViewLink": "https://drive/view/existing-id"}
    existing = {"id": "existing-id", "name": "Houndstooth Job Tracker.xlsx",
                "mimeType": tracker.XLSX_MIME}

    url, warning = tracker.upload_tracker(service, "folder1", out_path, "Houndstooth Job Tracker.xlsx", existing)

    service.files.return_value.update.assert_called_once()
    kwargs = service.files.return_value.update.call_args.kwargs
    assert kwargs["fileId"] == "existing-id"
    assert url == "https://drive/view/existing-id"
    assert warning is None
    service.files.return_value.create.assert_not_called()


def test_upload_existing_native_sheet_creates_new_and_warns(tmp_path):
    out_path = tmp_path / "t.xlsx"
    Workbook().save(out_path)
    service = MagicMock()
    service.files.return_value.create.return_value.execute.return_value = {
        "id": "new-id", "webViewLink": "https://drive/view/new-id"}
    existing = {"id": "old-sheet-id", "name": "Houndstooth Job Tracker",
                "mimeType": tracker.GSHEET_MIME}

    url, warning = tracker.upload_tracker(service, "folder1", out_path, "Houndstooth Job Tracker.xlsx", existing)

    service.files.return_value.create.assert_called_once()
    assert url == "https://drive/view/new-id"
    assert warning is not None
    assert "old-sheet-id" in warning
    service.files.return_value.update.assert_not_called()


def test_upload_no_existing_file_creates(tmp_path):
    out_path = tmp_path / "t.xlsx"
    Workbook().save(out_path)
    service = MagicMock()
    service.files.return_value.create.return_value.execute.return_value = {
        "id": "brand-new-id", "webViewLink": "https://drive/view/brand-new-id"}

    url, warning = tracker.upload_tracker(service, "folder1", out_path, "Houndstooth Job Tracker.xlsx", None)

    service.files.return_value.create.assert_called_once()
    create_kwargs = service.files.return_value.create.call_args.kwargs
    assert create_kwargs["body"]["name"] == "Houndstooth Job Tracker.xlsx"
    assert create_kwargs["body"]["parents"] == ["folder1"]
    assert url == "https://drive/view/brand-new-id"
    assert warning is None


def test_find_existing_tracker_returns_most_recently_modified(tmp_path):
    service = MagicMock()
    service.files.return_value.list.return_value.execute.return_value = {
        "files": [
            {"id": "newest", "name": "Houndstooth Job Tracker.xlsx", "mimeType": tracker.XLSX_MIME,
             "modifiedTime": "2026-08-05T00:00:00Z"},
        ]
    }
    result = tracker.find_existing_tracker(service, "folder1")
    assert result["id"] == "newest"


def test_find_existing_tracker_returns_none_when_absent(tmp_path):
    service = MagicMock()
    service.files.return_value.list.return_value.execute.return_value = {"files": []}
    result = tracker.find_existing_tracker(service, "folder1")
    assert result is None


# ---------------- dedup_key helper ----------------

def test_dedup_key_uses_stripped_url_when_present():
    assert tracker.dedup_key("  https://x/1  ", "Title", "Co") == "https://x/1"


def test_dedup_key_falls_back_to_nolink_when_url_empty():
    assert tracker.dedup_key("", "Title", "Co") == "NOLINK::Title::Co"
    assert tracker.dedup_key(None, "Title", "Co") == "NOLINK::Title::Co"
